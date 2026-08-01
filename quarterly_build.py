#!/usr/bin/env python3
"""
Auto-update the QUARTERLY geographic ATM tabs (Geography · State / District) in index.html.

Scrapes RBI's "State Wise and Region Wise Deployment of ATMs" page, downloads any new
quarterly workbook, parses its Region/State/District sheets, MERGES into the EMBEDDED_Q
block already in index.html, and rewrites EMBEDDED_Q + EMBEDDED_Q_STAMP. No change -> exit 0.
Runs alongside the monthly build_dashboard.py.
"""
import os, re, sys, io, json, time, datetime
LIST_URL="https://www.rbi.org.in/Scripts/StateRegionATMView.aspx"
DOC_RE=re.compile(r'https://rbidocs\.rbi\.org\.in/[^"\'<> ]+?\.XLSX', re.I)
HTML_PATH=os.path.join(os.path.dirname(__file__),"index.html")
UA=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36")

GROUPS_Q=["Public Sector Banks","Private Sector Banks","Foreign Banks","Small Finance Banks",
 "District Central Co-Operative Banks","Local Area Banks","Regional Rural Banks",
 "State Co-Operative Banks","Urban Co-Operative Banks","White Label ATM Operators (WLAOs)"]
GSHORT={"Public Sector Banks":"Public Sector","Private Sector Banks":"Private Sector","Foreign Banks":"Foreign",
 "Small Finance Banks":"Small Finance","District Central Co-Operative Banks":"District Central Co-op",
 "Local Area Banks":"Local Area","Regional Rural Banks":"Regional Rural","State Co-Operative Banks":"State Co-op",
 "Urban Co-Operative Banks":"Urban Co-op","White Label ATM Operators (WLAOs)":"White Label ATM"}
# same canonical names + rebrand renames as the monthly dashboard, so the same institution
# never appears twice under a spelling/rebrand variant
_ALIAS={"CITY UNION BANK":"CITY UNION BANK LTD","IDBI LTD":"IDBI BANK LTD",
 "JAMMU AND KASHMIR BANK":"JAMMU AND KASHMIR BANK LTD","SBM BANK INDIA":"SBM BANK INDIA LTD",
 "BANDHAN BANK":"BANDHAN BANK LTD","DHANALAKSHMI BANK LTD":"DHANALAXMI BANK LTD",
 "DBS BANK":"DBS INDIA BANK LTD","CATHOLIC SYRIAN BANK LTD":"CSB BANK LTD",
 "IDFC BANK LTD":"IDFC FIRST BANK LTD","RATNAKAR BANK LTD":"RBL BANK LTD",
 "DEVELOPMENT CREDIT BANK":"DCB BANK LTD","HONGKONG AND SHANGHAI BKG CORPN":"HSBC LTD",
 "AMERICAN EXPRESS":"AMERICAN EXPRESS BANKING CORPORATION",
 "NORTH EAST SMALL FINANCE BANK LTD":"SLICE SMALL FINANCE BANK LTD"}
# amalgamations: absorbed bank's rows folded (summed) into its acquirer
MERGE_Q={"FINCARE SMALL FINANCE BANK LTD":"AU SMALL FINANCE BANK LTD",
 "VIJAYA BANK":"BANK OF BARODA","DENA BANK":"BANK OF BARODA","ANDHRA BANK":"UNION BANK OF INDIA",
 "CORPORATION BANK":"UNION BANK OF INDIA","ORIENTAL BANK OF COMMERCE":"PUNJAB NATIONAL BANK",
 "UNITED BANK OF INDIA":"PUNJAB NATIONAL BANK","SYNDICATE BANK":"CANARA BANK","ALLAHABAD BANK":"INDIAN BANK",
 "THE LAXMI VILAS BANK LTD":"DBS INDIA BANK LTD"}
_M={'jan':'Jan','feb':'Feb','mar':'Mar','apr':'Apr','may':'May','jun':'Jun','jul':'Jul','aug':'Aug','sep':'Sep','oct':'Oct','nov':'Nov','dec':'Dec'}
_MO=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

def canon(n):
    s=re.sub(r"\s+"," ",str(n).strip()); u=re.sub(r"\.","",s.upper())
    u=re.sub(r"\bLIMITED\b","LTD",u)
    u=re.sub(r"CO[\s\-]*OPERATIVE","CO-OPERATIVE",u)   # unify co-operative / cooperative / co operative
    u=re.sub(r"\s+"," ",u).strip()
    return _ALIAS.get(u,u)

def fold_geo(d):
    """Fold amalgamated banks into their acquirer (sum region + state vectors); drop old names."""
    for key in ("region","state"):
        rows=d[key]
        for tgt,acq in MERGE_Q.items():
            if tgt in rows:
                tv=rows.pop(tgt)
                if acq in rows: rows[acq]=[a+b for a,b in zip(rows[acq],tv)]
                else: rows[acq]=tv
    for tgt,acq in MERGE_Q.items():
        if tgt in d["rgroups"]:
            d["rgroups"].setdefault(acq, d["rgroups"].pop(tgt))
            d["rgroups"].pop(tgt,None)
    return d
def num(x): return x if isinstance(x,(int,float)) else 0
def is_grp(s): return s in GROUPS_Q
def is_total(s): return 'total' in s.lower()
def q_sort_key(q): a,b=q.split(); return (int(b),_MO.index(a)+1)
def quarter_label(txt):
    m=re.search(r"(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)[A-Z]*[\s\-]*'?(\d{2,4})", str(txt), re.I)
    if not m: return None
    y=int(m.group(2)); y=2000+y if y<100 else y
    return f"{_M[m.group(1)[:3].lower()]} {y}"
def _find(wb,kw):
    for sn in wb.sheetnames:
        if kw in sn.lower(): return sn
    return None

def parse_quarterly(data):
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    rsn=_find(wb,"region"); ssn=_find(wb,"state"); dsn=_find(wb,"district")
    if not (rsn and ssn and dsn): return None
    region={}; rgroups={}; cur=None
    rows=list(wb[rsn].iter_rows(values_only=True))
    hi=next((i for i,r in enumerate(rows) if r and any('metropolitan' in str(c).lower() for c in r if c)),3)
    for r in rows[hi+1:]:
        c=r[1] if len(r)>1 else None
        if c is None: continue
        s=str(c).strip()
        if is_grp(s): cur=GSHORT[s]; continue
        if is_total(s) or s.lower()=='banks': continue
        if not any(isinstance(r[j],(int,float)) for j in range(2,6)): continue
        nm=canon(s); region[nm]=[num(r[2]),num(r[3]),num(r[4]),num(r[5])]; rgroups[nm]=cur or "Other"
    srows=list(wb[ssn].iter_rows(values_only=True))
    shi=next((i for i,r in enumerate(srows) if r and any(str(c).strip().upper() in ('BANK NAME','BANKS','BANK') for c in r if c)),3)
    states=[str(c).strip() for c in srows[shi][2:] if c is not None and str(c).strip() and 'grand total' not in str(c).lower()]
    state={}; cur=None
    for r in srows[shi+1:]:
        c=r[1] if len(r)>1 else None
        if c is None: continue
        s=str(c).strip()
        if is_grp(s): cur=GSHORT[s]; continue
        if is_total(s) or s.lower() in ('bank name','banks'): continue
        vals=[num(r[2+i]) for i in range(len(states))]
        if not any(vals): continue
        state[canon(s)]=vals
    drows=list(wb[dsn].iter_rows(values_only=True))
    dhi=next((i for i,r in enumerate(drows) if r and any('district' in str(c).lower() for c in r if c)),3)
    district={}; cur=None
    for r in drows[dhi+1:]:
        st=r[1] if len(r)>1 else None; di=r[2] if len(r)>2 else None
        if st is not None and str(st).strip():
            if is_total(str(st)) or str(st).strip().upper()=='NAME OF STATE': continue
            cur=str(st).strip()
        if di is None or not str(di).strip() or is_total(str(di)): continue
        district.setdefault(cur,[]).append([str(di).strip(),num(r[3]),num(r[4]),num(r[5]),num(r[6])])
    if not region or not state: return None
    return fold_geo({"region":region,"rgroups":rgroups,"state":state,"states":states,"district":district})

# ---------- fetch (requests, then Playwright fallback) ----------
def get_links():
    import requests
    r=requests.get(LIST_URL,headers={"User-Agent":UA},timeout=60)
    return sorted(set(DOC_RE.findall(r.text)))
def fetch_files(urls):
    out={}
    try:
        import requests
        s=requests.Session(); s.headers.update({"User-Agent":UA,"Referer":LIST_URL})
        try: s.get(LIST_URL,timeout=60)
        except Exception: pass
        for u in urls:
            try:
                r=s.get(u,timeout=90)
                if r.content[:2]==b'PK': out[u]=r.content
            except Exception: pass
        if len(out)==len(urls): return out
    except Exception as e: print("requests path:",e)
    missing=[u for u in urls if u not in out]
    if missing:
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                br=p.chromium.launch(args=["--no-sandbox"]); ctx=br.new_context(user_agent=UA); pg=ctx.new_page()
                pg.goto(LIST_URL,wait_until="networkidle",timeout=90000); pg.wait_for_timeout(2500)
                for u in missing:
                    try:
                        rr=ctx.request.get(u,headers={"Referer":LIST_URL},timeout=90000); b=rr.body()
                        if b[:2]==b'PK': out[u]=b
                    except Exception as e: print("  dl fail",u[-30:],e)
                br.close()
        except Exception as e: print("playwright path:",e)
    return out

# ---------- splice ----------
def read_embedded_q(html):
    m=re.search(r'const EMBEDDED_Q = (\{.*?\});\nconst EMBEDDED_Q_STAMP', html, re.S)
    if not m: return None
    try: return json.loads(m.group(1))
    except Exception: return None
def splice_q(html, emb, stamp):
    js="const EMBEDDED_Q = "+json.dumps(emb,separators=(",",":"))+";\nconst EMBEDDED_Q_STAMP = "+str(stamp)+";"
    html=re.sub(r'const EMBEDDED_Q = \{.*?\};\nconst EMBEDDED_Q_STAMP = \d+;', lambda _:js, html, count=1, flags=re.S)
    # also bump the main stamp so re-hosted file is treated as newest
    cur=int(re.search(r'const EMBEDDED_STAMP = (\d+)',html).group(1))
    html=re.sub(r'const EMBEDDED_STAMP = \d+','const EMBEDDED_STAMP = '+str(max(cur+1,int(time.time()*1000))),html,count=1)
    return html

def main():
    if not os.path.exists(HTML_PATH): print("ERROR: index.html missing"); sys.exit(1)
    html=open(HTML_PATH,encoding="utf-8").read()
    emb=read_embedded_q(html)
    if emb is None: print("ERROR: EMBEDDED_Q block not found in index.html"); sys.exit(1)
    have=set(emb.get("quarterOrder",[]))
    try: urls=get_links()
    except Exception as e: print("ERROR: listing page:",e); sys.exit(1)
    if not urls: print("ERROR: no quarterly XLSX links found."); sys.exit(1)
    # only fetch files whose label is new
    want={}
    for u in urls:
        ql=quarter_label(u.split("/")[-1])
        if ql and ql not in have: want[ql]=u
    if not want:
        print("No new quarter on the RBI page — nothing to do."); return
    print("New quarter file(s):", ", ".join(sorted(want)))
    files=fetch_files(list(want.values()))
    master=emb["states"]; midx={s:i for i,s in enumerate(master)}
    added=[]
    for u,b in files.items():
        d=parse_quarterly(b)
        if not d: continue
        ql=quarter_label(u.split("/")[-1])
        if not ql: continue
        # extend master with any new states
        for s in d["states"]:
            if s not in midx: midx[s]=len(master); master.append(s)
        st={}
        for bank,v in d["state"].items():
            idx={s:i for i,s in enumerate(d["states"])}
            st[bank]=[ v[idx[s]] if s in idx else 0 for s in master ]
        dist={k:v for k,v in d["district"].items() if k in midx}
        emb["quarters"][ql]={"region":d["region"],"rgroups":d["rgroups"],"state":st,"district":dist}
        added.append(ql)
    if not added:
        print("Downloaded but parsed no new quarter."); return
    # pad older quarters' state vectors if master grew
    for q,qd in emb["quarters"].items():
        for bank,v in qd["state"].items():
            if len(v)<len(master): qd["state"][bank]=v+[0]*(len(master)-len(v))
    emb["states"]=master
    emb["quarterOrder"]=sorted(set(list(have)+added), key=q_sort_key)
    stamp=int(time.time()*1000)
    open(HTML_PATH,"w",encoding="utf-8").write(splice_q(html,emb,stamp))
    print("Added quarter(s):",", ".join(added),"| total quarters:",len(emb["quarterOrder"]))

if __name__=="__main__":
    main()
